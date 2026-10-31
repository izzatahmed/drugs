"""Import a completed RenalRounds review sheet (.xlsx) into reviews/<file>.json.

Usage:  python3 import_review.py RenalRounds_review_sheet_returned.xlsx
Then rebuild: apply_reviews.py runs in the build after the release scripts.

- Drug sign-off rows with a decision, reviewer name and date are recorded.
- Reference columns (BNF, Lexidrugs, UpToDate, MIMS), TGA pregnancy category and the JFDA
  product check are recorded as reviewer statements; no reference text is stored.
- "Change needed" dose rows are listed in the JSON under "corrections" for a maintainer to apply
  against the cited source. Corrections are never applied automatically.
"""
import json, sys, datetime
from pathlib import Path
sys.dont_write_bytecode = True
from openpyxl import load_workbook

R = Path(__file__).parent
DECISION = {'Approve': 'approved', 'Approve after changes': 'changes_requested', 'Not approved': 'not_approved', 'Not my area': 'not_my_area'}
REF = {'Agrees': 'agrees', 'Differs': 'differs', 'Not listed': 'not_listed', 'Not checked': 'not_checked'}
JFDA = {'Matches': 'matches', 'Differs': 'differs', 'Not registered in Jordan': 'not_registered', 'Not checked': 'not_checked'}
TGA = {'A', 'B1', 'B2', 'B3', 'C', 'D', 'X', 'Exempt'}

def iso(v):
    if isinstance(v, (datetime.date, datetime.datetime)): return v.strftime('%Y-%m-%d')
    s = str(v or '').strip()
    try: return datetime.date.fromisoformat(s[:10]).isoformat()
    except ValueError: return None

def main(path):
    wb = load_workbook(path, data_only=True)
    so, dr = wb['Drug sign-off'], wb['Dose rows']
    head = {so.cell(1, j).value: j for j in range(1, so.max_column + 1)}
    need = ['Entry ID', 'Your decision', 'Reviewer name', 'Review date']
    missing = [h for h in need if h not in head]
    if missing: raise SystemExit(f'Not a RenalRounds review sheet (missing columns: {missing})')
    col = lambda r, name: so.cell(r, head[name]).value if name in head else None
    drugs, problems = {}, []
    for r in range(2, so.max_row + 1):
        did = col(r, 'Entry ID')
        if not did: continue
        entry = {}
        dec = col(r, 'Your decision')
        if dec:
            who, when = (col(r, 'Reviewer name') or '').strip(), iso(col(r, 'Review date'))
            if dec not in DECISION: problems.append(f'{did}: unknown decision "{dec}"'); continue
            if not who or not when: problems.append(f'{did}: decision without reviewer name or valid date'); continue
            entry['review'] = dict(decision=DECISION[dec], reviewer=who, date=when, comment=col(r, 'Comment') or '')
        refs = {}
        for key, name in [('bnf', 'BNF'), ('lexidrugs', 'Lexidrugs'), ('uptodate', 'UpToDate'), ('mims', 'MIMS (Australia)')]:
            v = col(r, name)
            if v and v in REF: refs[key] = REF[v]
            elif v: problems.append(f'{did}: unknown {name} value "{v}"')
        if refs:
            entry['references'] = refs; entry['referenceNote'] = col(r, 'Reference notes') or ''
            entry['referenceBy'] = (col(r, 'Reviewer name') or '').strip(); entry['referenceDate'] = iso(col(r, 'Review date'))
        t = col(r, 'TGA pregnancy category')
        if t:
            t = str(t).strip().upper().replace('EXEMPT', 'Exempt')
            if t in TGA: entry['tga'] = t
            else: problems.append(f'{did}: unknown TGA category "{t}"')
        j = col(r, 'JFDA product check')
        if j:
            if j in JFDA: entry['jfda'] = dict(status=JFDA[j], note=col(r, 'JFDA notes') or '', by=(col(r, 'Reviewer name') or '').strip())
            else: problems.append(f'{did}: unknown JFDA value "{j}"')
        if entry: drugs[did] = entry
    dh = {dr.cell(1, j).value: j for j in range(1, dr.max_column + 1)}
    corrections = []
    for r in range(2, dr.max_row + 1):
        if dr.cell(r, dh['Your decision']).value in ('Change needed', 'Remove'):
            corrections.append({k: dr.cell(r, dh[k]).value for k in ['Entry ID', 'Setting', 'Table / indication', 'Kidney function / context', 'Dose or statement shown in app', 'Your decision', 'Corrected wording', 'Comment']})
    out = R/'reviews'; out.mkdir(exist_ok=True)
    target = out/(Path(path).stem + '.json')
    target.write_text(json.dumps(dict(sourceFile=Path(path).name, importedAt=datetime.date.today().isoformat(), drugs=drugs, corrections=corrections, problems=problems), ensure_ascii=False, indent=1))
    approved = sum(1 for e in drugs.values() if e.get('review', {}).get('decision') == 'approved')
    print(json.dumps(dict(file=str(target), drugsWithEntries=len(drugs), approved=approved, corrections=len(corrections), problems=problems[:20]), ensure_ascii=False, indent=1))

if __name__ == '__main__':
    if len(sys.argv) != 2: raise SystemExit(__doc__)
    main(sys.argv[1])
