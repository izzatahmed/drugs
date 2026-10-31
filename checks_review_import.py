"""Round-trip test: fill a copy of the review sheet, import it, apply it to a scratch database.
Usage: python3 checks_review_import.py /path/RenalRounds_review_sheet.xlsx"""
import json, shutil, subprocess, sys, tempfile, datetime
from pathlib import Path
sys.dont_write_bytecode = True
from openpyxl import load_workbook
R = Path(__file__).parent; sheet = Path(sys.argv[1])
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp); w = R
    for name in ['import_review.py', 'apply_reviews.py']: shutil.copy(w/name, tmp/name)
    (tmp/'dist').mkdir(); shutil.copy(w/'dist'/'database.json', tmp/'dist'/'database.json')
    wb = load_workbook(sheet); so = wb['Drug sign-off']; head = {so.cell(1, j).value: j for j in range(1, so.max_column + 1)}
    rows = {so.cell(r, head['Entry ID']).value: r for r in range(2, so.max_row + 1)}
    def put(did, **kw):
        for k, v in kw.items(): so.cell(rows[did], head[k]).value = v
    put('vancomycin', **{'Your decision': 'Approve', 'Reviewer name': 'Dr Test', 'Review date': datetime.date(2026, 11, 2), 'BNF': 'Agrees', 'Lexidrugs': 'Agrees', 'TGA pregnancy category': 'B2', 'JFDA product check': 'Matches'})
    put('gentamicin', **{'Your decision': 'Approve after changes', 'Reviewer name': 'Dr Test', 'Review date': '2026-11-02', 'UpToDate': 'Differs', 'Reference notes': 'HD dose range'})
    put('apixaban', **{'Your decision': 'Approve', 'Reviewer name': ''})      # must be rejected: no reviewer
    filled = tmp/'filled.xlsx'; wb.save(filled)
    out = subprocess.run([sys.executable, 'import_review.py', str(filled)], cwd=tmp, capture_output=True, text=True); assert out.returncode == 0, out.stderr
    subprocess.run([sys.executable, 'apply_reviews.py'], cwd=tmp, check=True, capture_output=True)
    db = {d['id']: d for d in json.loads((tmp/'dist'/'database.json').read_text())['drugs']}
    v, g, a = db['vancomycin'], db['gentamicin'], db['apixaban']
    assert v['review']['status'] == 'clinically_reviewed' and v['review']['clinicalReviewer'] == 'Dr Test' and v['review']['clinicalApprovedAt'] == '2026-11-02'
    assert v['referenceChecks']['bnf']['status'] == 'agrees' and v['referenceChecks']['bnf']['checkedBy'] == 'Dr Test'
    assert v['pregnancyLactation']['tgaCategory'] == 'B2' and v['referenceChecks']['jfda']['status'] == 'matches'
    assert g['review']['status'] == 'changes_requested' and g['review']['clinicalApprovedAt'] is None and g['referenceChecks']['uptodate']['status'] == 'differs'
    assert a['review']['status'] == 'clinical_review_pending', 'approval without a reviewer name must be ignored'
    rep = json.loads(next((tmp/'reviews').glob('*.json')).read_text()); assert any('apixaban' in p for p in rep['problems'])
print('PASS: review sheet import/apply round trip (approval, changes requested, reference checks, TGA category, JFDA check, rejection of unsigned approval).')
